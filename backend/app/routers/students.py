"""学生名单与管理：``/api/students``。

* GET    ``/api/students``            列出启用学生（按学号升序）。
* POST   ``/api/students``            新增学生（学号唯一，重复 409）。
* PATCH  ``/api/students/{id}``       改名 / 改学号。
* DELETE ``/api/students/{id}``       停用（active=0，保留历史作文关联）。
* POST   ``/api/students/import``     批量导入：学号已存在视为改名更新。
"""

from __future__ import annotations

import csv
import io
from typing import Annotated, Any

from fastapi import APIRouter, Depends, File, Query, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.auth import require_auth
from app.db import get_session
from app.models import Student, utcnow_iso
from app.render.templates import natural_text_key
from app.schemas import (
    ApiError,
    Envelope,
    StudentCreate,
    StudentImportRequest,
    StudentOut,
    StudentUpdate,
    envelope,
)

router = APIRouter(prefix="/api/students", tags=["students"])

SessionDep = Annotated[AsyncSession, Depends(get_session)]
AuthDep = Annotated[str, Depends(require_auth)]

_HEADER_ALIASES: dict[str, tuple[str, ...]] = {
    "student_no": ("学号", "学生学号", "学号编号", "studentno", "student_no", "no"),
    "name": ("姓名", "学生姓名", "名字", "name", "studentname"),
}


def _normalized_header(value: object) -> str:
    return "".join(str(value or "").strip().lower().split())


def _find_columns(headers: list[object]) -> tuple[int | None, int | None]:
    normalized = [_normalized_header(value) for value in headers]
    no_index = next(
        (index for index, value in enumerate(normalized) if value in {_normalized_header(alias) for alias in _HEADER_ALIASES["student_no"]}),
        None,
    )
    name_index = next(
        (index for index, value in enumerate(normalized) if value in {_normalized_header(alias) for alias in _HEADER_ALIASES["name"]}),
        None,
    )
    return no_index, name_index


def _decode_csv(data: bytes) -> list[list[object]]:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = data.decode(encoding)
            return [list(row) for row in csv.reader(io.StringIO(text))]
        except UnicodeDecodeError:
            continue
    raise ValueError("CSV 文件编码无法识别，请另存为 UTF-8")


def _read_roster_file(filename: str, data: bytes) -> list[list[object]]:
    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "csv":
        return _decode_csv(data)
    if suffix != "xlsx":
        raise ValueError("只支持 .xlsx 或 .csv 文件")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sheet = workbook.active
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        workbook.close()
        return rows
    except Exception as exc:
        raise ValueError("Excel 文件无法读取，请使用模板填写后重试") from exc


def _parse_roster_rows(rows: list[list[object]]) -> tuple[list[dict[str, str]], list[dict[str, Any]]]:
    if not rows:
        return [], [{"row": 1, "message": "文件没有数据"}]
    no_index, name_index = _find_columns(rows[0])
    errors: list[dict[str, Any]] = []
    if no_index is None or name_index is None:
        return [], [{"row": 1, "message": "表头必须包含学号和姓名字段"}]

    parsed: list[dict[str, str]] = []
    seen: set[str] = set()
    for row_number, row in enumerate(rows[1:], start=2):
        student_no = str(row[no_index] if no_index < len(row) and row[no_index] is not None else "").strip()
        name = str(row[name_index] if name_index < len(row) and row[name_index] is not None else "").strip()
        if not student_no and not name:
            continue
        if not student_no or not name:
            errors.append({"row": row_number, "message": "学号和姓名不能为空"})
            continue
        if len(student_no) > 50 or len(name) > 100:
            errors.append({"row": row_number, "message": "学号最多 50 字，姓名最多 100 字"})
            continue
        if student_no in seen:
            errors.append({"row": row_number, "message": f"学号 {student_no} 在文件内重复"})
            continue
        seen.add(student_no)
        parsed.append({"student_no": student_no, "name": name})
    return parsed, errors


async def _preview_rows(
    session: AsyncSession, parsed: list[dict[str, str]], errors: list[dict[str, Any]]
) -> dict[str, Any]:
    existing = {
        student.student_no: student
        for student in (
            await session.execute(select(Student).where(Student.student_no.in_([item["student_no"] for item in parsed])))
        ).scalars().all()
    }
    rows = [
        {
            **item,
            "action": "更新" if item["student_no"] in existing else "新增",
        }
        for item in parsed
    ]
    return {
        "rows": rows,
        "errors": errors,
        "valid_count": len(parsed),
        "invalid_count": len(errors),
        "created_count": sum(1 for item in rows if item["action"] == "新增"),
        "updated_count": sum(1 for item in rows if item["action"] == "更新"),
    }


@router.get("/import-template")
async def download_import_template(
    _auth: AuthDep, file_format: Annotated[str, Query(alias="format")] = "xlsx"
) -> Response:
    """下载学生名单模板，字段固定为学号、姓名。"""
    if file_format not in {"xlsx", "csv"}:
        raise ApiError("模板格式只支持 xlsx 或 csv", code=400, status_code=400)
    if file_format == "csv":
        data = "学号,姓名\n20230101,张三\n".encode("utf-8-sig")
        return Response(data, media_type="text/csv; charset=utf-8", headers={"Content-Disposition": 'attachment; filename="students-template.csv"'})
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["学号", "姓名"])
    sheet.append(["20230101", "张三"])
    output = io.BytesIO()
    workbook.save(output)
    return Response(output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": 'attachment; filename="students-template.xlsx"'})


@router.post("/import-file", response_model=Envelope[dict[str, Any]])
async def preview_import_file(
    file: Annotated[UploadFile, File(...)], session: SessionDep, _auth: AuthDep
) -> dict[str, Any]:
    """解析名单文件并返回预览；解析阶段绝不写数据库。"""
    data = await file.read()
    if len(data) > 5 * 1024 * 1024:
        raise ApiError("名单文件不能超过 5 MB", code=400, status_code=400)
    try:
        parsed, errors = _parse_roster_rows(_read_roster_file(file.filename or "", data))
    except ValueError as exc:
        raise ApiError(str(exc), code=400, status_code=400) from exc
    return envelope(await _preview_rows(session, parsed, errors))


@router.post("/import-file/confirm", response_model=Envelope[dict[str, Any]])
async def confirm_import_file(
    payload: StudentImportRequest, session: SessionDep, _auth: AuthDep
) -> dict[str, Any]:
    """确认导入预览中的有效行；复用批量导入的事务语义，整批成功或失败。"""
    seen: set[str] = set()
    for item in payload.students:
        if item.student_no in seen:
            raise ApiError(f"导入清单内学号 {item.student_no} 重复", code=400, status_code=400)
        seen.add(item.student_no)

    existing = {
        student.student_no: student
        for student in (
            await session.execute(select(Student).where(Student.student_no.in_(seen)))
        ).scalars().all()
    }
    created = updated = 0
    for item in payload.students:
        current = existing.get(item.student_no)
        if current is None:
            session.add(
                Student(
                    student_no=item.student_no,
                    name=item.name,
                    active=1,
                    created_at=utcnow_iso(),
                )
            )
            created += 1
        elif current.active != 1 or current.name != item.name:
            current.name = item.name
            current.active = 1
            updated += 1
    await session.commit()
    return envelope({"created": created, "updated": updated}, message="学生名单已导入")


@router.get("", response_model=Envelope[list[StudentOut]])
async def list_students(session: SessionDep, _auth: AuthDep) -> dict[str, Any]:
    """列出启用中的学生（按学号升序）。"""
    stmt = select(Student).where(Student.active == 1).order_by(Student.student_no.asc())
    students = sorted(
        (await session.execute(stmt)).scalars().all(),
        key=lambda student: natural_text_key(student.student_no),
    )
    return envelope([StudentOut.model_validate(student) for student in students])


async def _ensure_no_taken(session: AsyncSession, student_no: str, exclude_id: int | None) -> None:
    """学号唯一性校验（exclude_id 供更新时排除自身）。"""
    stmt = select(Student).where(Student.student_no == student_no)
    if exclude_id is not None:
        stmt = stmt.where(Student.id != exclude_id)
    existing = (await session.execute(stmt)).scalar_one_or_none()
    if existing is not None:
        raise ApiError(f"学号 {student_no} 已存在", code=409, status_code=409)


@router.post("", status_code=201, response_model=Envelope[StudentOut])
async def create_student(
    payload: StudentCreate, session: SessionDep, _auth: AuthDep
) -> dict[str, Any]:
    """新增一名学生。"""
    await _ensure_no_taken(session, payload.student_no, exclude_id=None)
    student = Student(
        student_no=payload.student_no,
        name=payload.name,
        active=1,
        created_at=utcnow_iso(),
    )
    session.add(student)
    await session.commit()
    await session.refresh(student)
    return envelope(StudentOut.model_validate(student))


@router.patch("/{student_id}", response_model=Envelope[StudentOut])
async def update_student(
    student_id: int, payload: StudentUpdate, session: SessionDep, _auth: AuthDep
) -> dict[str, Any]:
    """更新学生姓名 / 学号。"""
    student = await session.get(Student, student_id)
    if student is None or student.active != 1:
        raise ApiError("学生不存在", code=404, status_code=404)

    if payload.student_no is not None and payload.student_no != student.student_no:
        await _ensure_no_taken(session, payload.student_no, exclude_id=student.id)
        student.student_no = payload.student_no
    if payload.name is not None:
        student.name = payload.name

    await session.commit()
    await session.refresh(student)
    return envelope(StudentOut.model_validate(student))


@router.delete("/{student_id}", response_model=Envelope[dict[str, Any]])
async def deactivate_student(
    student_id: int, session: SessionDep, _auth: AuthDep
) -> dict[str, Any]:
    """停用学生（active=0，保留历史作文关联，不从名单删除）。"""
    student = await session.get(Student, student_id)
    if student is None:
        raise ApiError("学生不存在", code=404, status_code=404)

    student.active = 0
    await session.commit()
    return envelope({"id": student_id}, message="学生已停用")


@router.post("/import", response_model=Envelope[dict[str, Any]])
async def import_students(
    payload: StudentImportRequest, session: SessionDep, _auth: AuthDep
) -> dict[str, Any]:
    """批量导入学生：学号已存在则更新姓名，否则新增。

    Raises:
        ApiError: 400 学号在本次导入清单内重复。
    """
    seen: set[str] = set()
    for item in payload.students:
        if item.student_no in seen:
            raise ApiError(f"导入清单内学号 {item.student_no} 重复", code=400, status_code=400)
        seen.add(item.student_no)

    existing = {
        s.student_no: s
        for s in (
            await session.execute(
                select(Student).where(Student.student_no.in_(seen))  # noqa: S608
            )
        ).scalars().all()
    }

    created = updated = 0
    for item in payload.students:
        current = existing.get(item.student_no)
        if current is not None:
            if current.active != 1 or current.name != item.name:
                current.name = item.name
                current.active = 1  # 重新导入即恢复启用
                updated += 1
        else:
            session.add(
                Student(
                    student_no=item.student_no,
                    name=item.name,
                    active=1,
                    created_at=utcnow_iso(),
                )
            )
            created += 1

    await session.commit()
    return envelope({"created": created, "updated": updated})
